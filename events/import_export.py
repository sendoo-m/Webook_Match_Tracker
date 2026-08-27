import csv
import io
from datetime import date, datetime, time
from zoneinfo import ZoneInfo

import openpyxl

from matches.models import Venue

from .models import Category, Event
from .services import attach_default_checklist_items

RIYADH_TZ = ZoneInfo("Asia/Riyadh")

EXPORT_FIELDS = [
    "slug",
    "category",
    "venue",
    "title_ar",
    "title_en",
    "description_ar",
    "description_en",
    "event_date",
    "start_time",
    "end_time",
    "gates_open_time",
    "sale_starts_at",
]


def _event_row_values(event):
    return [
        event.slug,
        event.category.name_en,
        event.venue.name_en if event.venue else "",
        event.title_ar,
        event.title_en,
        event.description_ar,
        event.description_en,
        event.event_date.isoformat() if event.event_date else "",
        event.start_time.strftime("%H:%M") if event.start_time else "",
        event.end_time.strftime("%H:%M") if event.end_time else "",
        event.gates_open_time.strftime("%H:%M") if event.gates_open_time else "",
        event.sale_starts_at.astimezone(RIYADH_TZ).strftime("%Y-%m-%d %H:%M") if event.sale_starts_at else "",
    ]


def export_events_csv(queryset):
    buffer = io.StringIO()
    # BOM so Excel opens the CSV as UTF-8 instead of guessing a codepage that
    # mangles Arabic - xlsx is still the recommended format for Arabic text.
    buffer.write("﻿")
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_FIELDS)

    for event in queryset.select_related("category", "venue"):
        writer.writerow(_event_row_values(event))
    return buffer.getvalue()


def export_events_xlsx(queryset):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Events"
    ws.append(EXPORT_FIELDS)

    for event in queryset.select_related("category", "venue"):
        ws.append(_event_row_values(event))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_import_template_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Events"
    ws.append(EXPORT_FIELDS)

    reference = wb.create_sheet("Reference")
    reference.append(["Categories (use exactly as shown)"])
    categories = list(
        Category.objects.filter(is_active=True).order_by("group__sort_order", "sort_order").values_list("name_en", flat=True)
    )
    for name in categories:
        reference.append([name])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class ImportResult:
    def __init__(self):
        self.created = 0
        self.updated = 0
        self.errors = []
        self.skipped = []

    @property
    def has_errors(self):
        return bool(self.errors)

    @property
    def total_ok(self):
        return self.created + self.updated


def _clean_str(value):
    if value is None:
        return ""
    return str(value).strip()


def _find_category(name):
    name = _clean_str(name)
    if not name:
        return None
    return Category.objects.filter(name_en=name).first() or Category.objects.filter(name_ar=name).first()


def _to_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = str(value).strip()
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return None


def _to_time(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    value = str(value).strip()
    if not value:
        return None
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p"):
        try:
            return datetime.strptime(value, fmt).time()
        except ValueError:
            continue
    return None


def _to_datetime(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=RIYADH_TZ)
    value = str(value).strip()
    if not value:
        return None
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            naive = datetime.strptime(value, fmt)
            return naive.replace(tzinfo=RIYADH_TZ)
        except ValueError:
            continue
    return None


def _import_row(row, row_number, result, category_filter_id=None):
    slug = _clean_str(row.get("slug"))
    if not slug:
        result.errors.append((row_number, "Missing slug"))
        return

    category = _find_category(row.get("category"))
    if not category:
        result.errors.append((row_number, f"Unknown category: {row.get('category', '')!r}"))
        return

    if category_filter_id and category.id != category_filter_id:
        result.skipped.append((row_number, f"Skipped - {category.name_en} not in selected filter"))
        return

    venue = None
    venue_name = _clean_str(row.get("venue"))
    if venue_name:
        venue, _ = Venue.objects.get_or_create(
            name_en=venue_name, defaults={"name_ar": venue_name}
        )

    try:
        event, created = Event.objects.update_or_create(
            slug=slug,
            defaults={
                "category": category,
                "venue": venue,
                "title_ar": _clean_str(row.get("title_ar")),
                "title_en": _clean_str(row.get("title_en")),
                "description_ar": _clean_str(row.get("description_ar")),
                "description_en": _clean_str(row.get("description_en")),
                "event_date": _to_date(row.get("event_date")),
                "start_time": _to_time(row.get("start_time")),
                "end_time": _to_time(row.get("end_time")),
                "gates_open_time": _to_time(row.get("gates_open_time")),
                "sale_starts_at": _to_datetime(row.get("sale_starts_at")),
            },
        )
    except Exception as exc:
        result.errors.append((row_number, str(exc)))
        return

    if created:
        attach_default_checklist_items(event)
        result.created += 1
    else:
        result.updated += 1


def import_events_csv(file_obj, category_filter_id=None):
    """Row-by-row update_or_create keyed on slug."""
    result = ImportResult()
    text_stream = io.TextIOWrapper(file_obj, encoding="utf-8-sig")
    reader = csv.DictReader(text_stream)

    for row_number, row in enumerate(reader, start=2):
        _import_row(row, row_number, result, category_filter_id)

    return result


def import_events_xlsx(file_obj, category_filter_id=None):
    result = ImportResult()
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)

    try:
        header = [_clean_str(h) for h in next(rows)]
    except StopIteration:
        return result

    for row_number, values in enumerate(rows, start=2):
        if all(v is None for v in values):
            continue
        row = {header[i]: values[i] for i in range(min(len(header), len(values)))}
        _import_row(row, row_number, result, category_filter_id)

    return result


def import_events_file(file_obj, filename, category_filter_id=None):
    if filename.lower().endswith(".xlsx"):
        return import_events_xlsx(file_obj, category_filter_id)
    return import_events_csv(file_obj, category_filter_id)
