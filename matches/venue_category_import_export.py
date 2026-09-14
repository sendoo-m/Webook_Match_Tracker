# matches/venue_category_import_export.py
#
# Excel import/export for VenueSeatingCategory, scoped to one (venue, club)
# pair per file - a category list can differ between two clubs sharing the
# same physical venue, so the file only ever represents one club's list at
# one venue, never a cross-venue/cross-club dump. Same openpyxl-based
# pattern as matches/import_export.py's match import/export.

import io

import openpyxl

from matches.import_export import ImportResult
from matches.models import VenueSeatingCategory

CATEGORY_EXPORT_FIELDS = ["code", "seat_count"]


def _clean_str(value):
    if value is None:
        return ""
    return str(value).strip()


def _to_int(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(value)
    value = str(value).strip()
    if not value:
        return None
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


def export_venue_categories_xlsx(venue, club):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Categories"
    ws.append(CATEGORY_EXPORT_FIELDS)

    categories = VenueSeatingCategory.objects.filter(venue=venue, club=club).order_by("sort_order", "code")
    for category in categories:
        ws.append([category.code, category.seat_count if category.seat_count is not None else ""])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_venue_category_import_template_xlsx(venue, club):
    """Same as export - a template pre-filled with the venue/club's
    existing category codes if any exist yet, or just the header row for
    a brand-new venue/club pair. Reusing export as the template avoids
    keeping two near-identical builders in sync."""
    return export_venue_categories_xlsx(venue, club)


def import_venue_categories_xlsx(file_obj, venue, club):
    """Upserts VenueSeatingCategory rows for this (venue, club) pair only -
    keyed by code. Rows for other venues/clubs cannot be created through
    this path; the venue/club are fixed for the whole file, not read from
    it."""
    result = ImportResult()
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)

    try:
        header = [_clean_str(h).lower() for h in next(rows)]
    except StopIteration:
        return result

    # Accepts both this app's own export column names (code/seat_count)
    # and the names commonly used in ticket-listing spreadsheets exported
    # from elsewhere ("Ticket Name"/"Total Capacity") - "Price"/any other
    # column is simply not looked up, so it's ignored without erroring.
    code_aliases = ("code", "ticket name", "category", "category code")
    seat_count_aliases = ("seat_count", "total capacity", "capacity", "seat count")

    code_idx = next((header.index(name) for name in code_aliases if name in header), None)
    if code_idx is None:
        result.errors.append((1, 'Missing required column "code" (or "Ticket Name").'))
        return result
    seat_count_idx = next((header.index(name) for name in seat_count_aliases if name in header), None)

    for row_number, values in enumerate(rows, start=2):
        if all(v is None for v in values):
            continue

        code = _clean_str(values[code_idx]) if code_idx < len(values) else ""
        if not code:
            result.errors.append((row_number, "Missing category code."))
            continue

        seat_count = None
        if seat_count_idx is not None and seat_count_idx < len(values):
            seat_count = _to_int(values[seat_count_idx])

        category, created = VenueSeatingCategory.objects.update_or_create(
            venue=venue,
            club=club,
            code=code,
            defaults={"seat_count": seat_count},
        )
        if created:
            result.created += 1
        else:
            result.updated += 1

    return result
