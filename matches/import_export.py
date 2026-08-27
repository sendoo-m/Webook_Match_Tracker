import csv
import io
from datetime import date, datetime, time
from zoneinfo import ZoneInfo

import openpyxl

from checklists.services import attach_default_checklist_items
from matches.models import Club, Competition, Match, Venue

RIYADH_TZ = ZoneInfo("Asia/Riyadh")

EXPORT_FIELDS = [
    "slug",
    "competition",
    "home_club",
    "away_club",
    "venue",
    "round_number",
    "title_ar",
    "title_en",
    "description_ar",
    "description_en",
    "event_date",
    "match_start_time",
    "match_end_time",
    "gates_open_time",
    "sale_starts_at",
]


def _match_row_values(match):
    return [
        match.slug,
        match.competition.name_en,
        match.home_club.name_en,
        match.away_club.name_en,
        match.venue.name_en if match.venue else "",
        match.round_number if match.round_number is not None else "",
        match.title_ar,
        match.title_en,
        match.description_ar,
        match.description_en,
        match.event_date.isoformat() if match.event_date else "",
        match.match_start_time.strftime("%H:%M") if match.match_start_time else "",
        match.match_end_time.strftime("%H:%M") if match.match_end_time else "",
        match.gates_open_time.strftime("%H:%M") if match.gates_open_time else "",
        match.sale_starts_at.astimezone(RIYADH_TZ).strftime("%Y-%m-%d %H:%M") if match.sale_starts_at else "",
    ]


def export_matches_csv(queryset):
    buffer = io.StringIO()
    # BOM so Excel opens the CSV as UTF-8 instead of guessing a codepage that
    # mangles Arabic - xlsx is still the recommended format for Arabic text.
    buffer.write("﻿")
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_FIELDS)

    for match in queryset.select_related("competition", "home_club", "away_club", "venue"):
        writer.writerow(_match_row_values(match))
    return buffer.getvalue()


def export_matches_xlsx(queryset):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matches"
    ws.append(EXPORT_FIELDS)

    for match in queryset.select_related("competition", "home_club", "away_club", "venue"):
        ws.append(_match_row_values(match))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


SPL_REPORT_FIELDS = [
    "Matchweek",
    "Match Date",
    "Home Team",
    "Away Team",
    "Stadium",
    "Release Date",
    "Actual Release Date",
    "Match Readiness %",
    "Ticketing Plan Status",
    "KV / Webook Images",
    "Webook Readiness",
    "SPL Complimentary Tickets",
    "Comments",
    "Coordinator (Webook)",
]


def _spl_report_row_values(row):
    match = row["match"]
    return [
        match.round_number if match.round_number is not None else "",
        match.event_date.isoformat() if match.event_date else "",
        match.home_club.name_en,
        match.away_club.name_en,
        match.venue.name_en if match.venue else "",
        match.sale_starts_at.astimezone(RIYADH_TZ).strftime("%Y-%m-%d %H:%M") if match.sale_starts_at else "",
        match.actual_release_at.astimezone(RIYADH_TZ).strftime("%Y-%m-%d %H:%M") if match.actual_release_at else "",
        f"{row['readiness_percent']}%",
        "Approved" if match.ticketing_plan_approved else "Not Approved",
        "Ready" if row["kv_ready"] else "Not Ready",
        "Ready" if row["webook_ready"] else "Not Ready",
        "Sent" if match.spl_tickets_sent else "Not Sent",
        match.spl_comments,
        (match.home_club.owner.get_full_name() or match.home_club.owner.username) if match.home_club.owner_id else "",
    ]


def export_spl_report_xlsx(rows):
    """rows: iterable of dicts shaped like operations.views.helpers.build_spl_report_row's output."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SPL Report"
    ws.append(SPL_REPORT_FIELDS)

    for row in rows:
        ws.append(_spl_report_row_values(row))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_import_template_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matches"
    ws.append(EXPORT_FIELDS)

    reference = wb.create_sheet("Reference")
    reference.append(["Competitions (use exactly as shown)", "Clubs (use exactly as shown)"])
    competitions = list(Competition.objects.filter(is_active=True).order_by("sort_order").values_list("name_en", flat=True))
    clubs = list(Club.objects.filter(is_active=True).order_by("name_ar").values_list("name_en", flat=True))
    for i in range(max(len(competitions), len(clubs))):
        reference.append([
            competitions[i] if i < len(competitions) else "",
            clubs[i] if i < len(clubs) else "",
        ])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class ImportResult:
    def __init__(self):
        self.created = 0
        self.updated = 0
        self.errors = []
        self.skipped = []

    @property
    def has_errors(self):
        return bool(self.errors)

    @property
    def total_ok(self):
        return self.created + self.updated


def _clean_str(value):
    if value is None:
        return ""
    return str(value).strip()


def _find_club(name):
    name = _clean_str(name)
    if not name:
        return None
    return Club.objects.filter(name_en=name).first() or Club.objects.filter(name_ar=name).first()


def _find_competition(name):
    name = _clean_str(name)
    if not name:
        return None
    return Competition.objects.filter(name_en=name).first() or Competition.objects.filter(name_ar=name).first()


def _to_int(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(value)
    value = str(value).strip()
    if not value:
        return None
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


def _to_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = str(value).strip()
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return None


def _to_time(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    value = str(value).strip()
    if not value:
        return None
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p"):
        try:
            return datetime.strptime(value, fmt).time()
        except ValueError:
            continue
    return None


def _to_datetime(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=RIYADH_TZ)
    value = str(value).strip()
    if not value:
        return None
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            naive = datetime.strptime(value, fmt)
            return naive.replace(tzinfo=RIYADH_TZ)
        except ValueError:
            continue
    return None


def _import_row(row, row_number, result, competition_filter_id=None):
    slug = _clean_str(row.get("slug"))
    if not slug:
        result.errors.append((row_number, "Missing slug"))
        return

    competition = _find_competition(row.get("competition"))
    if not competition:
        result.errors.append((row_number, f"Unknown competition: {row.get('competition', '')!r}"))
        return

    if competition_filter_id and competition.id != competition_filter_id:
        result.skipped.append((row_number, f"Skipped - {competition.name_en} not in selected filter"))
        return

    home_club = _find_club(row.get("home_club"))
    away_club = _find_club(row.get("away_club"))
    if not home_club or not away_club:
        result.errors.append(
            (row_number, f"Unknown club(s): home={row.get('home_club', '')!r}, away={row.get('away_club', '')!r}")
        )
        return
    if home_club == away_club:
        result.errors.append((row_number, "Home and away clubs must be different"))
        return

    venue = None
    venue_name = _clean_str(row.get("venue"))
    if venue_name:
        venue, _ = Venue.objects.get_or_create(
            name_en=venue_name, defaults={"name_ar": venue_name}
        )

    try:
        match, created = Match.objects.update_or_create(
            slug=slug,
            defaults={
                "competition": competition,
                "home_club": home_club,
                "away_club": away_club,
                "venue": venue,
                "round_number": _to_int(row.get("round_number")),
                "title_ar": _clean_str(row.get("title_ar")),
                "title_en": _clean_str(row.get("title_en")),
                "description_ar": _clean_str(row.get("description_ar")),
                "description_en": _clean_str(row.get("description_en")),
                "event_date": _to_date(row.get("event_date")),
                "match_start_time": _to_time(row.get("match_start_time")),
                "match_end_time": _to_time(row.get("match_end_time")),
                "gates_open_time": _to_time(row.get("gates_open_time")),
                "sale_starts_at": _to_datetime(row.get("sale_starts_at")),
            },
        )
    except Exception as exc:
        result.errors.append((row_number, str(exc)))
        return

    if created:
        attach_default_checklist_items(match)
        result.created += 1
    else:
        result.updated += 1


def import_matches_csv(file_obj, competition_filter_id=None):
    """Row-by-row update_or_create keyed on slug."""
    result = ImportResult()
    text_stream = io.TextIOWrapper(file_obj, encoding="utf-8-sig")
    reader = csv.DictReader(text_stream)

    for row_number, row in enumerate(reader, start=2):
        _import_row(row, row_number, result, competition_filter_id)

    return result


def import_matches_xlsx(file_obj, competition_filter_id=None):
    result = ImportResult()
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)

    try:
        header = [_clean_str(h) for h in next(rows)]
    except StopIteration:
        return result

    for row_number, values in enumerate(rows, start=2):
        if all(v is None for v in values):
            continue
        row = {header[i]: values[i] for i in range(min(len(header), len(values)))}
        _import_row(row, row_number, result, competition_filter_id)

    return result


def import_matches_file(file_obj, filename, competition_filter_id=None):
    if filename.lower().endswith(".xlsx"):
        return import_matches_xlsx(file_obj, competition_filter_id)
    return import_matches_csv(file_obj, competition_filter_id)
