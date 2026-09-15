# clubs/pricing_import_export.py
#
# Excel import for a club's per-category pricing plan submission - one
# file always maps to exactly one match's current home club/venue, never
# picked from the file itself. Same openpyxl-based header-alias pattern
# as matches/venue_category_import_export.py.

import io

import openpyxl

from matches.import_export import ImportResult
from matches.models import VenueSeatingCategory


def _clean_str(value):
    if value is None:
        return ""
    return str(value).strip()


def _to_decimal(value):
    if value is None or value == "":
        return None
    try:
        return round(float(value), 2)
    except (ValueError, TypeError):
        return None


def build_plan_category_price_template_xlsx(categories):
    """A ready-to-fill spreadsheet for the "Or Upload an Excel File" import
    below: one row per this match's own active seating categories, code
    pre-filled, price left blank for the club to type in and re-upload
    as-is - avoids the club having to type category codes by hand or get
    one wrong."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append(["code", "price"])

    for category in categories:
        ws.append([category.code, ""])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def parse_plan_category_prices_xlsx(file_obj, venue, club):
    """Reads (category code, price) pairs from an uploaded Excel file and
    resolves each code against this venue/club's active
    VenueSeatingCategory rows. Does NOT write anything - returns
    (prices_by_category_id, result) so the caller can create the
    ClubPricingPlan version and price rows in one transaction only if the
    whole file is clean, rather than partially importing a bad file."""
    result = ImportResult()
    prices_by_category_id = {}

    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)

    try:
        header = [_clean_str(h).lower() for h in next(rows)]
    except StopIteration:
        return prices_by_category_id, result

    code_aliases = ("code", "ticket name", "category", "category code")
    price_aliases = ("price", "ticket price")

    code_idx = next((header.index(name) for name in code_aliases if name in header), None)
    price_idx = next((header.index(name) for name in price_aliases if name in header), None)
    if code_idx is None:
        result.errors.append((1, 'Missing required column "code" (or "Ticket Name").'))
        return prices_by_category_id, result
    if price_idx is None:
        result.errors.append((1, 'Missing required column "price".'))
        return prices_by_category_id, result

    categories_by_code = {
        c.code.strip().lower(): c
        for c in VenueSeatingCategory.objects.filter(venue=venue, club=club, is_active=True)
    }

    for row_number, values in enumerate(rows, start=2):
        if all(v is None for v in values):
            continue

        code = _clean_str(values[code_idx]) if code_idx < len(values) else ""
        if not code:
            result.errors.append((row_number, "Missing category code."))
            continue

        category = categories_by_code.get(code.lower())
        if category is None:
            result.errors.append((row_number, f'Unknown or inactive category code "{code}" for this venue/club.'))
            continue

        price = _to_decimal(values[price_idx]) if price_idx < len(values) else None
        if price is None:
            result.errors.append((row_number, f'Missing or invalid price for category "{code}".'))
            continue

        prices_by_category_id[category.id] = price
        result.created += 1

    return prices_by_category_id, result
